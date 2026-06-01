#!/usr/bin/env python3
"""bank transaction CSV → XLSX budget tracker with outstanding authorisations."""
import argparse
import re
import subprocess
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BUDGET  = 6250.00
SAVINGS = 3000.00
OUTPUT_DIR = Path(__file__).parent / "output" / "transaction_v2"
COL_NAMES = ["Date", "Amount", "Type", "Description"]

MONTHS_ABBR = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}


def load_transactions(csv_path: Path) -> pd.DataFrame:
    raw = pd.read_csv(csv_path, header=None, names=["Date", "_amt", "Description"])
    raw["_amt"] = pd.to_numeric(
        raw["_amt"].astype(str).str.replace('"', "").str.strip(),
        errors="coerce",
    )
    df = pd.DataFrame()
    df["Date"] = pd.to_datetime(raw["Date"], dayfirst=True)
    df["Amount"] = raw["_amt"].abs()
    df["Type"] = raw["_amt"].apply(lambda x: "Credit" if x > 0 else "Expense")
    df["Description"] = raw["Description"]
    return df


def _get_outstanding_text() -> str:
    placeholder = (
        "Paste your Outstanding Authorisations text here, replacing this message.\n"
        "Click Import when done, or Skip to continue without outstanding items.\n"
    )
    result = subprocess.run(
        ["zenity", "--text-info", "--editable",
         "--title=Paste Outstanding Authorisations",
         "--width=900", "--height=500",
         "--ok-label=Import",
         "--cancel-label=Skip"],
        input=placeholder,
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        return ""
    text = result.stdout.strip()
    if "replacing this message" in text:
        return ""
    return text


def parse_outstanding(text: str, year: int) -> pd.DataFrame:
    DATE_PAT = re.compile(r"^(\d{1,2})\s+([A-Z]{3})\s*$")
    AMOUNT_PAT = re.compile(r"\$([\d,]+\.\d{2})")

    lines = [l.strip() for l in text.splitlines()]
    rows = []
    i = 0

    while i < len(lines):
        m = DATE_PAT.match(lines[i])
        if m:
            day, mon = int(m.group(1)), m.group(2)
            month = MONTHS_ABBR.get(mon)
            if not month:
                i += 1
                continue
            try:
                tx_date = date(year, month, day)
            except ValueError:
                i += 1
                continue

            i += 1
            while i < len(lines) and (not lines[i] or lines[i] == "POS AUTHORISATION"):
                i += 1

            if i < len(lines) and lines[i]:
                desc_line = lines[i]
                amounts = AMOUNT_PAT.findall(desc_line)
                if amounts:
                    amt_str = amounts[0]
                    desc = desc_line[:desc_line.index("$" + amt_str)].strip()
                    rows.append({
                        "Date": pd.Timestamp(tx_date),
                        "Amount": float(amt_str.replace(",", "")),
                        "Type": "Outstanding",
                        "Description": desc,
                    })
                i += 1
            continue
        i += 1

    if not rows:
        return pd.DataFrame(columns=COL_NAMES)
    return pd.DataFrame(rows)


def _post_process(xlsx_path: Path, df: pd.DataFrame) -> None:
    SUMMARY_ROWS = 8  # title + budget + confirmed + outstanding + credits + net spent + remaining + savings
    ncols = len(COL_NAMES)

    wb = load_workbook(xlsx_path)
    ws = wb.active
    ws.insert_rows(1, SUMMARY_ROWS)

    data_header = SUMMARY_ROWS + 1
    data_start  = data_header + 1
    data_end    = len(df) + SUMMARY_ROWS + 1

    ws.auto_filter.ref = f"A{data_header}:{get_column_letter(ncols)}{data_end}"
    ws.freeze_panes = f"A{data_start}"

    amt_col  = get_column_letter(COL_NAMES.index("Amount") + 1)
    type_col = get_column_letter(COL_NAMES.index("Type") + 1)
    rng_amt  = f"${amt_col}${data_start}:${amt_col}${data_end}"
    rng_type = f"${type_col}${data_start}:${type_col}${data_end}"

    thin = Side(style="thin")
    box  = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("G1:H1")
    ws["G1"] = "Summary"
    ws["G1"].font = Font(name="Calibri", bold=True)

    summary_rows = [
        ("Budget",             BUDGET),
        ("Confirmed Expenses", f'=SUMIF({rng_type},"Expense",{rng_amt})'),
        ("Outstanding",        f'=SUMIF({rng_type},"Outstanding",{rng_amt})'),
        ("Total Credits",      f'=SUMIF({rng_type},"Credit",{rng_amt})'),
        ("Net Spent",          "=H3+H4-H5"),
        ("Remaining",          "=H2-H6"),
        ("Savings",            f"={SAVINGS}+MIN(0,H7)"),
    ]
    for i, (label, value) in enumerate(summary_rows, start=2):
        ws[f"G{i}"] = label
        ws[f"G{i}"].font = Font(name="Calibri")
        ws[f"H{i}"] = value
        ws[f"H{i}"].number_format = "$#,##0.00"
        ws[f"H{i}"].font = Font(name="Calibri")

    # Remaining: green ≥0, red <0
    ws.conditional_formatting.add("H7", CellIsRule(
        operator="greaterThanOrEqual", formula=["0"],
        fill=PatternFill("solid", fgColor="C6EFCE"),
    ))
    ws.conditional_formatting.add("H7", CellIsRule(
        operator="lessThan", formula=["0"],
        fill=PatternFill("solid", fgColor="FFB3B3"),
    ))
    # Savings: 3-tier
    for cell_ref in ("H8", "K4"):
        ws.conditional_formatting.add(cell_ref, CellIsRule(
            operator="lessThan", formula=["0"],
            fill=PatternFill("solid", fgColor="FFB3B3"),
        ))
        ws.conditional_formatting.add(cell_ref, CellIsRule(
            operator="lessThan", formula=[str(SAVINGS)],
            fill=PatternFill("solid", fgColor="FFE0B2"),
        ))
        ws.conditional_formatting.add(cell_ref, CellIsRule(
            operator="greaterThanOrEqual", formula=[str(SAVINGS)],
            fill=PatternFill("solid", fgColor="C6EFCE"),
        ))
    # Projected Remaining: green ≥0, red <0
    ws.conditional_formatting.add("K3", CellIsRule(
        operator="greaterThanOrEqual", formula=["0"],
        fill=PatternFill("solid", fgColor="C6EFCE"),
    ))
    ws.conditional_formatting.add("K3", CellIsRule(
        operator="lessThan", formula=["0"],
        fill=PatternFill("solid", fgColor="FFB3B3"),
    ))

    for row in ws["G1:H8"]:
        for cell in row:
            cell.border = box

    # Projected box in J:K (col I is a spacer)
    ws.merge_cells("J1:K1")
    ws["J1"] = "Projected"
    ws["J1"].font = Font(name="Calibri", bold=True)
    projected_rows = [
        ("Tentative", f'=SUMIF({rng_type},"Tentative",{rng_amt})'),
        ("Remaining", "=H7-K2"),
        ("Savings",   f"={SAVINGS}+MIN(0,K3)"),
    ]
    for i, (label, value) in enumerate(projected_rows, start=2):
        ws[f"J{i}"] = label
        ws[f"J{i}"].font = Font(name="Calibri")
        ws[f"K{i}"] = value
        ws[f"K{i}"].number_format = "$#,##0.00"
        ws[f"K{i}"].font = Font(name="Calibri")
    for row in ws["J1:K4"]:
        for cell in row:
            cell.border = box

    amt_col_idx  = COL_NAMES.index("Amount") + 1
    type_col_idx = COL_NAMES.index("Type") + 1
    highlight_to = type_col_idx

    expense_fills = [
        (1000, PatternFill("solid", fgColor="FFB3B3")),
        (600,  PatternFill("solid", fgColor="FFCCBC")),
        (300,  PatternFill("solid", fgColor="FFE0B2")),
        (80,   PatternFill("solid", fgColor="FFF2CC")),
    ]
    credit_fill    = PatternFill("solid", fgColor="C6EFCE")
    tentative_fill = PatternFill("solid", fgColor="E8DAEF")  # light lavender

    for r in range(data_header, data_end + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Calibri", bold=cell.font.bold if cell.font else False)

    for r in range(data_start, data_end + 1):
        ws.cell(row=r, column=amt_col_idx).number_format = "$#,##0.00"
        type_val = ws.cell(row=r, column=type_col_idx).value
        amt_val  = ws.cell(row=r, column=amt_col_idx).value

        if type_val == "Credit":
            for c in range(1, highlight_to + 1):
                ws.cell(row=r, column=c).fill = credit_fill
        elif type_val == "Tentative":
            for c in range(1, highlight_to + 1):
                ws.cell(row=r, column=c).fill = tentative_fill
                ws.cell(row=r, column=c).font = Font(name="Calibri", bold=True)
        elif isinstance(amt_val, (int, float)):
            for threshold, fill in expense_fills:
                if amt_val >= threshold:
                    for c in range(1, highlight_to + 1):
                        ws.cell(row=r, column=c).fill = fill
                        ws.cell(row=r, column=c).font = Font(name="Calibri", bold=True)
                    break

    for col in ws.iter_cols():
        col_letter = get_column_letter(col[0].column)
        max_len = max(
            (
                len(f"${cell.value:,.2f}") if isinstance(cell.value, (int, float))
                else len(str(cell.value)) if cell.value is not None and not str(cell.value).startswith("=")
                else 0
                for cell in col
            ),
            default=8,
        )
        ws.column_dimensions[col_letter].width = max_len + 2

    ws.column_dimensions["I"].width = 2  # spacer between Summary and Projected
    proj_label_width = max(len(label) for label, _ in projected_rows) + 2
    proj_value_width = len(f"${BUDGET:,.2f}") + 2
    ws.column_dimensions["J"].width = proj_label_width
    ws.column_dimensions["K"].width = proj_value_width

    wb.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert bank transaction CSV to XLSX with outstanding authorisations.")
    parser.add_argument("csv_path", type=Path, nargs="?")
    args = parser.parse_args()

    csv_path = args.csv_path
    if csv_path is None:
        result = subprocess.run(
            ["zenity", "--file-selection", "--title=Select bank CSV", "--file-filter=*.csv"],
            capture_output=True, text=True,
        )
        chosen = result.stdout.strip()
        if not chosen:
            raise SystemExit("No file selected.")
        csv_path = Path(chosen)

    OUTPUT_DIR.mkdir(exist_ok=True)
    df = load_transactions(csv_path)

    outstanding_text = _get_outstanding_text()
    if outstanding_text:
        year = df["Date"].max().year
        outstanding_df = parse_outstanding(outstanding_text, year)
        if not outstanding_df.empty:
            df = pd.concat([df, outstanding_df], ignore_index=True)

    blank_tentative = pd.DataFrame([
        {"Date": None, "Amount": None, "Type": "Tentative", "Description": ""}
        for _ in range(5)
    ])
    df = pd.concat([df, blank_tentative], ignore_index=True)
    df = df.sort_values("Date", ascending=False, na_position="last").reset_index(drop=True)

    xlsx_path = OUTPUT_DIR / csv_path.with_suffix(".xlsx").name

    if xlsx_path.exists():
        answer = input(f"{xlsx_path.name} already exists. [o]verwrite / [v]ersion up? ").strip().lower()
        if answer == "v":
            stem = xlsx_path.stem
            n = 2
            while xlsx_path.exists():
                xlsx_path = OUTPUT_DIR / f"{stem}_v{n}.xlsx"
                n += 1
        elif answer != "o":
            raise SystemExit("Cancelled.")

    df.to_excel(xlsx_path, index=False, engine="openpyxl")
    _post_process(xlsx_path, df)
    print(xlsx_path)
    subprocess.Popen(["xdg-open", str(xlsx_path)])


if __name__ == "__main__":
    main()
