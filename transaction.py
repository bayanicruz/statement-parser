#!/usr/bin/env python3
"""Transaction CSV → XLSX budget tracker."""
import argparse
import subprocess
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BUDGET  = 6250.00
SAVINGS = 3000.00
OUTPUT_DIR = Path(__file__).parent / "output" / "transactions"
COL_NAMES = ["Date", "Amount", "Type", "Description", "Category"]


def load_transactions(csv_path: Path) -> pd.DataFrame:
    raw = pd.read_csv(csv_path, header=None, names=["Date", "_amt", "Description"])  # headerless: Date, Amount, Description
    raw["_amt"] = pd.to_numeric(
        raw["_amt"].astype(str).str.replace('"', "").str.strip(),
        errors="coerce",
    )
    df = pd.DataFrame()
    df["Date"] = raw["Date"]
    df["Amount"] = raw["_amt"].abs()
    df["Type"] = raw["_amt"].apply(lambda x: "Credit" if x > 0 else "Expense")
    df["Description"] = raw["Description"]
    df["Category"] = ""
    return df


def _post_process(xlsx_path: Path, df: pd.DataFrame) -> None:
    SUMMARY_ROWS = 7  # title + budget + expenses + credits + net spent + remaining + savings
    ncols = len(COL_NAMES)

    wb = load_workbook(xlsx_path)
    ws = wb.active
    ws.insert_rows(1, SUMMARY_ROWS)

    data_header = SUMMARY_ROWS + 1
    data_start  = data_header + 1
    data_end    = len(df) + SUMMARY_ROWS + 1

    ws.auto_filter.ref = f"A{data_header}:{get_column_letter(ncols)}{data_end}"
    ws.freeze_panes = f"A{data_start}"

    amt_col  = get_column_letter(COL_NAMES.index("Amount") + 1)
    type_col = get_column_letter(COL_NAMES.index("Type") + 1)
    rng_amt  = f"${amt_col}${data_start}:${amt_col}${data_end}"
    rng_type = f"${type_col}${data_start}:${type_col}${data_end}"

    thin = Side(style="thin")
    box  = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("G1:H1")
    ws["G1"] = "Summary"
    ws["G1"].font = Font(name="Calibri", bold=True)

    summary_rows = [
        ("Budget",         BUDGET),
        ("Total Expenses", f'=SUMIF({rng_type},"Expense",{rng_amt})'),
        ("Total Credits",  f'=SUMIF({rng_type},"Credit",{rng_amt})'),
        ("Net Spent",      "=H3-H4"),
        ("Remaining",      "=H2-H5"),
        ("Savings",        f"={SAVINGS}+MIN(0,H6)"),
    ]
    for i, (label, value) in enumerate(summary_rows, start=2):
        ws[f"G{i}"] = label
        ws[f"G{i}"].font = Font(name="Calibri")
        ws[f"H{i}"] = value
        ws[f"H{i}"].number_format = "$#,##0.00"
        ws[f"H{i}"].font = Font(name="Calibri")

    ws.conditional_formatting.add("H6", CellIsRule(
        operator="greaterThanOrEqual", formula=["0"],
        fill=PatternFill("solid", fgColor="C6EFCE"),
    ))
    ws.conditional_formatting.add("H6", CellIsRule(
        operator="lessThan", formula=["0"],
        fill=PatternFill("solid", fgColor="FFB3B3"),
    ))
    ws.conditional_formatting.add("H7", CellIsRule(
        operator="lessThan", formula=["0"],
        fill=PatternFill("solid", fgColor="FFB3B3"),
    ))
    ws.conditional_formatting.add("H7", CellIsRule(
        operator="lessThan", formula=[str(SAVINGS)],
        fill=PatternFill("solid", fgColor="FFE0B2"),
    ))
    ws.conditional_formatting.add("H7", CellIsRule(
        operator="greaterThanOrEqual", formula=[str(SAVINGS)],
        fill=PatternFill("solid", fgColor="C6EFCE"),
    ))

    for row in ws["G1:H7"]:
        for cell in row:
            cell.border = box

    # Expense highlight tiers (cols A–C: Date, Amount, Type)
    amt_col_idx = COL_NAMES.index("Amount") + 1
    type_col_idx = COL_NAMES.index("Type") + 1
    highlight_to = type_col_idx  # Date, Amount, Type

    expense_fills = [
        (1000, PatternFill("solid", fgColor="FFB3B3")),  # red
        (600,  PatternFill("solid", fgColor="FFCCBC")),  # salmon
        (300,  PatternFill("solid", fgColor="FFE0B2")),  # orange
        (80,   PatternFill("solid", fgColor="FFF2CC")),  # yellow
    ]
    credit_fill = PatternFill("solid", fgColor="C6EFCE")  # green

    # Normalise all data fonts to Calibri
    for r in range(data_header, data_end + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Calibri", bold=cell.font.bold if cell.font else False)

    for r in range(data_start, data_end + 1):
        ws.cell(row=r, column=amt_col_idx).number_format = "$#,##0.00"
        type_val = ws.cell(row=r, column=type_col_idx).value
        amt_val = ws.cell(row=r, column=amt_col_idx).value

        if type_val == "Credit":
            for c in range(1, highlight_to + 1):
                ws.cell(row=r, column=c).fill = credit_fill
        elif isinstance(amt_val, (int, float)):
            for threshold, fill in expense_fills:
                if amt_val >= threshold:
                    for c in range(1, highlight_to + 1):
                        ws.cell(row=r, column=c).fill = fill
                        ws.cell(row=r, column=c).font = Font(name="Calibri", bold=True)
                    break

    for col in ws.iter_cols():
        max_len = max(
            (len(str(cell.value)) for cell in col
             if cell.value is not None and not str(cell.value).startswith("=")),
            default=8,
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    wb.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert transaction CSV to XLSX budget tracker.")
    parser.add_argument("csv_path", type=Path, nargs="?")
    args = parser.parse_args()

    csv_path = args.csv_path
    if csv_path is None:
        result = subprocess.run(
            ["zenity", "--file-selection", "--title=Select CSV", "--file-filter=*.csv"],
            capture_output=True, text=True,
        )
        chosen = result.stdout.strip()
        if not chosen:
            raise SystemExit("No file selected.")
        csv_path = Path(chosen)

    OUTPUT_DIR.mkdir(exist_ok=True)
    df = load_transactions(csv_path)
    xlsx_path = OUTPUT_DIR / csv_path.with_suffix(".xlsx").name

    if xlsx_path.exists():
        answer = input(f"{xlsx_path.name} already exists. [o]verwrite / [v]ersion up? ").strip().lower()
        if answer == "v":
            stem = xlsx_path.stem
            n = 2
            while xlsx_path.exists():
                xlsx_path = OUTPUT_DIR / f"{stem}_v{n}.xlsx"
                n += 1
        elif answer != "o":
            raise SystemExit("Cancelled.")

    df.to_excel(xlsx_path, index=False, engine="openpyxl")
    _post_process(xlsx_path, df)
    print(xlsx_path)
    subprocess.Popen(["xdg-open", str(xlsx_path)])


if __name__ == "__main__":
    main()
