#!/usr/bin/env python3
"""Credit card statement PDF → XLSX extractor."""
import argparse
import re
import subprocess
from collections import defaultdict
from collections.abc import Iterator
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent / "output" / "statements"
DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
CARD_RE = re.compile(r"^\d{4}$")
CARD_MAX_OFFSET = 80  # max horizontal distance (pts) from last date to card-number word
COL_NAMES = ["Processed", "Transaction", "Used", "Transaction Details", "Amount ($A)", "Balance"]


def _parse_row(words: list[dict]) -> list[str] | None:
    """Parse one row of PDF words into a transaction record, or None if not a data row."""
    ws = sorted(words, key=lambda w: w["x0"])
    dates = [w["text"] for w in ws if DATE_RE.match(w["text"])]
    if not dates:
        return None

    money = [w for w in ws if w["text"].startswith("$")]
    first_dollar_x0 = money[0]["x0"] if money else float("inf")
    last_date_x0 = max(w["x0"] for w in ws if DATE_RE.match(w["text"]))

    def is_card(w: dict) -> bool:
        return bool(CARD_RE.match(w["text"])) and w["x0"] - last_date_x0 < CARD_MAX_OFFSET

    card = next((w["text"] for w in ws if is_card(w)), "")
    desc = " ".join(
        w["text"] for w in ws
        if not DATE_RE.match(w["text"])
        and not is_card(w)
        and not w["text"].startswith("$")
        and w["x0"] < first_dollar_x0
    )

    if len(money) >= 2:
        amt_x0, bal_x0 = money[-2]["x0"], money[-1]["x0"]
        cr = "".join(w["text"] for w in ws if not w["text"].startswith("$") and amt_x0 < w["x0"] < bal_x0)
        amount = money[-2]["text"] + cr
    else:
        amount = ""

    return [dates[0], dates[1] if len(dates) > 1 else "", card, desc, amount, money[-1]["text"] if money else ""]


def _page_transactions(page) -> Iterator[list[str]]:
    """Yield transaction rows from a single PDF page."""
    by_top: dict[int, list] = defaultdict(list)
    for w in page.extract_words():
        by_top[round(w["top"])].append(w)

    in_table = False
    for _, words in sorted(by_top.items()):
        texts = {w["text"] for w in words}
        if not in_table:
            in_table = "Processed" in texts
            continue
        if "Please" in texts and "refer" in texts:
            break
        if row := _parse_row(words):
            yield row


def extract(pdf_path: str | Path) -> pd.DataFrame:
    """Extract all transactions from a statement PDF into a DataFrame."""
    with pdfplumber.open(pdf_path) as pdf:
        rows = [row for page in pdf.pages for row in _page_transactions(page)]
    if not rows:
        raise ValueError(f"No transactions found in {pdf_path}")
    df = pd.DataFrame(rows, columns=COL_NAMES)

    cr_mask = df["Amount ($A)"].str.contains("CR", na=False)
    numeric = df["Amount ($A)"].str.replace(r"[$,CR]", "", regex=True)
    df.insert(df.columns.get_loc("Amount ($A)") + 1, "Credit ($A)",
              pd.to_numeric(numeric.where(cr_mask), errors="coerce").mul(-1))
    df["Amount ($A)"] = pd.to_numeric(numeric.where(~cr_mask), errors="coerce")

    return df


def _post_process(xlsx_path: Path, df: pd.DataFrame) -> None:
    """Add AutoFilter, AUD currency format, and per-card Summary box to the workbook."""
    used_values = sorted(v for v in df["Used"].dropna().unique() if v)
    SUMMARY_ROWS = 1 + len(used_values) + 2  # title + one row per card + subtotal (filtered) + total

    wb = load_workbook(xlsx_path)
    ws = wb.active
    ws.insert_rows(1, SUMMARY_ROWS)

    data_header = SUMMARY_ROWS + 1
    data_start = data_header + 1
    data_end = len(df) + SUMMARY_ROWS + 1
    ncols = len(df.columns)
    used_col = get_column_letter(df.columns.get_loc("Used") + 1)
    amt_col = get_column_letter(df.columns.get_loc("Amount ($A)") + 1)
    used_range = f"${used_col}${data_start}:${used_col}${data_end}"
    amt_range = f"${amt_col}${data_start}:${amt_col}${data_end}"

    ws.auto_filter.ref = f"A{data_header}:{get_column_letter(ncols)}{data_end}"

    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("I1:J1")
    ws["I1"] = "Summary"
    ws["I1"].font = Font(name="Calibri", bold=True)

    for i, card in enumerate(used_values, start=2):
        ws[f"I{i}"] = card
        ws[f"J{i}"] = f'=SUMIF({used_range},"{card}",{amt_range})'
        ws[f"J{i}"].number_format = "$#,##0.00"

    subtotal_row = SUMMARY_ROWS - 1
    ws[f"I{subtotal_row}"] = "Subtotal (Filtered):"
    ws[f"J{subtotal_row}"] = f"=SUBTOTAL(9,{amt_range})"
    ws[f"J{subtotal_row}"].number_format = "$#,##0.00"

    ws[f"I{SUMMARY_ROWS}"] = "Total:"
    ws[f"J{SUMMARY_ROWS}"] = f"=SUM({amt_range})"
    ws[f"J{SUMMARY_ROWS}"].number_format = "$#,##0.00"

    for row in ws[f"I1:J{SUMMARY_ROWS}"]:
        for cell in row:
            cell.border = box

    aud = "$#,##0.00"
    amt_col_idx = df.columns.get_loc("Amount ($A)") + 1
    expense_fills = [
        (1000, PatternFill("solid", fgColor="FFB3B3")),  # red
        (600,  PatternFill("solid", fgColor="FFCCBC")),  # salmon
        (300,  PatternFill("solid", fgColor="FFE0B2")),  # orange
        (80,   PatternFill("solid", fgColor="FFF2CC")),  # yellow
    ]

    for r in range(data_header, data_end + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Calibri", bold=cell.font.bold if cell.font else False)

    for r in range(data_start, data_end + 1):
        ws.cell(row=r, column=amt_col_idx).number_format = aud
        ws.cell(row=r, column=df.columns.get_loc("Credit ($A)") + 1).number_format = aud
        cell = ws.cell(row=r, column=amt_col_idx)
        if isinstance(cell.value, (int, float)):
            for threshold, fill in expense_fills:
                if cell.value >= threshold:
                    for c in range(1, amt_col_idx + 1):
                        ws.cell(row=r, column=c).fill = fill
                        ws.cell(row=r, column=c).font = Font(name="Calibri", bold=True)
                    break

    for col in ws.iter_cols():
        max_len = max(
            (len(str(cell.value)) for cell in col
             if cell.value is not None and not str(cell.value).startswith("=")),
            default=8,
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    wb.save(xlsx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract transactions from a credit card statement PDF.")
    parser.add_argument("pdf_path", type=Path, nargs="?")
    args = parser.parse_args()

    pdf_path = args.pdf_path
    if pdf_path is None:
        result = subprocess.run(
            ["zenity", "--file-selection", "--title=Select statement PDF", "--file-filter=*.pdf"],
            capture_output=True, text=True,
        )
        chosen = result.stdout.strip()
        if not chosen:
            raise SystemExit("No file selected.")
        pdf_path = Path(chosen)

    OUTPUT_DIR.mkdir(exist_ok=True)
    df = extract(pdf_path)
    xlsx_path = OUTPUT_DIR / pdf_path.with_suffix(".xlsx").name

    if xlsx_path.exists():
        answer = input(f"{xlsx_path.name} already exists. [o]verwrite / [v]ersion up? ").strip().lower()
        if answer == "v":
            stem = xlsx_path.stem
            n = 2
            while xlsx_path.exists():
                xlsx_path = OUTPUT_DIR / f"{stem}_v{n}.xlsx"
                n += 1
        elif answer != "o":
            raise SystemExit("Cancelled.")
    df.to_excel(xlsx_path, index=False, engine="openpyxl")
    _post_process(xlsx_path, df)
    print(xlsx_path)
    subprocess.Popen(["xdg-open", str(xlsx_path)])


if __name__ == "__main__":
    main()
